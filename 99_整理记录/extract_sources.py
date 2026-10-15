from pathlib import Path
import os, json, hashlib, re, zipfile
from xml.etree import ElementTree as ET
from concurrent.futures import ThreadPoolExecutor
import pypdfium2 as pdfium
from pypdf import PdfReader
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

ROOT = Path('\\\\?\\' + str(Path(__file__).resolve().parents[1]))
OUT = ROOT / '99_整理记录'
CACHE = OUT / '文本缓存'
CACHE.mkdir(parents=True, exist_ok=True)
DOCEXT = {'.pdf', '.docx', '.xlsx', '.pptx', '.doc', '.xls'}
files = []
for base, dirs, names in os.walk(ROOT):
    dirs[:] = [d for d in dirs if d not in {'99_整理记录', '.git', '.codex', '.agents', '00_替代与差异知识库'}]
    for name in names:
        p = Path(base)/name
        s = p.stat()
        files.append({'path':p.relative_to(ROOT).as_posix(),'bytes':s.st_size,'mtime_ns':s.st_mtime_ns})
(OUT/'整理前文件清单.json').write_text(json.dumps(files, ensure_ascii=False, indent=2), encoding='utf-8')
docs = [f for f in files if Path(f['path']).suffix.lower() in DOCEXT or (Path(f['path']).suffix.lower()=='.txt' and len(Path(f['path']).parts)<=2)]
results=[]
for i, f in enumerate(sorted(docs,key=lambda f:f['path']),1):
    p = ROOT / f['path']; ext=p.suffix.lower(); parts=[]; units=0
    rec={'id':f'S{i:03d}',**f,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
    try:
        if ext=='.pdf':
            pdf=pdfium.PdfDocument(str(p)); units=len(pdf)
            # Reference books/TRMs: index first pages only; not the focus of this pass.
            selected=range(units) if units<=350 else range(min(15,units))
            for j in selected:
                page=pdf[j]; tp=page.get_textpage(); t=tp.get_text_range(); parts.append(f'=== PDF页 {j+1} ===\n{t}')
                tp.close(); page.close()
            pdf.close()
            rec['coverage']='全文提取' if len(selected)==units else f'仅前{len(selected)}页；大型参考手册未精读'
        elif ext=='.docx':
            doc=Document(p); n=0
            for el in doc.element.body:
                if el.tag.endswith('}p'):
                    t=Paragraph(el,doc).text
                    if t.strip(): n+=1; parts.append(f'段落{n}: {t}')
                elif el.tag.endswith('}tbl'):
                    n+=1; parts.append(f'表{n}:')
                    for ri,row in enumerate(Table(el,doc).rows,1): parts.append(f'行{ri}: '+' | '.join(c.text.replace('\n',' / ') for c in row.cells))
            units=n; rec['coverage']='正文与表格提取；图片文字需另查'
        elif ext=='.xlsx':
            wb=load_workbook(p,data_only=True,read_only=False); units=len(wb.worksheets)
            for ws in wb.worksheets:
                parts.append(f'=== 工作表 {ws.title} ===')
                for row in ws:
                    vals=[f'{c.coordinate}={c.value}' for c in row if c.value is not None]
                    if vals: parts.append(' | '.join(vals))
            wb.close(); rec['coverage']='所有工作表非空单元格；图片需另查'
        elif ext=='.pptx':
            with zipfile.ZipFile(p) as z:
                slides=sorted([n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)],key=lambda x:int(re.search(r'slide(\d+)',x)[1]))
                units=len(slides)
                for si,name in enumerate(slides,1):
                    root=ET.fromstring(z.read(name)); ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
                    txt=[''.join(e.itertext()) for e in root.findall('.//a:t',ns)]
                    parts.append(f'=== 幻灯片 {si} ===\n'+'\n'.join(txt))
            rec['coverage']='全部幻灯片文本；图片需另查'
        elif ext=='.txt':
            raw=p.read_bytes()
            for enc in ['utf-8-sig','gb18030','utf-16']:
                try: parts=[raw.decode(enc)]; break
                except UnicodeError: pass
            units=1; rec['coverage']='全文'
        else: rec['coverage']='旧版二进制文档，未提取'
        text='\n\n'.join(parts)
        rec.update(units=units, chars=len(text), cache=f'文本缓存/{rec["id"]}.txt')
        (OUT/rec['cache']).write_text(f'来源: {f["path"]}\n来源编号: {rec["id"]}\n\n'+text,encoding='utf-8')
    except Exception as exc: rec['error']=str(exc)
    results.append(rec)
    if i%20==0: print(f'Extracted {i}/{len(docs)}',flush=True)
(OUT/'来源索引.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'files':len(files),'bytes':sum(f['bytes'] for f in files),'documents':len(docs),'errors':[r for r in results if 'error' in r]},ensure_ascii=False),flush=True)
